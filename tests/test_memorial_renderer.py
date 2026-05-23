from __future__ import annotations

import json
from pathlib import Path
import unittest

from app.services.context_builder import build_memorial_eletrico_v1_context
from app.services.memorial_renderer import (
    MemorialRenderError,
    assert_no_internal_markers_left,
    assert_no_jinja_left,
    has_docx_render_dependencies,
    inspect_docx_text,
    render_memorial_eletrico_v1,
    render_memorial_gas_natural_v1,
    render_memorial_glp_v2,
    render_memorial_telecom_v1,
)
from app.services.memorial_validator import (
    validate_memorial_eletrico_v1_context,
    validate_memorial_gas_natural_v1_context,
    validate_memorial_glp_v2_context,
    validate_memorial_telecom_v1_context,
)


ROOT = Path(__file__).resolve().parent.parent
FIXTURE_PATH = ROOT / "tests" / "fixtures" / "eletrico_sem_subestacao.json"
TELECOM_FIXTURE_PATH = ROOT / "tests" / "fixtures" / "telecom_base.json"
GAS_NATURAL_FIXTURE_PATH = ROOT / "tests" / "fixtures" / "gas_natural_base.json"


def load_fixture() -> dict:
    with FIXTURE_PATH.open("r", encoding="utf-8") as file:
        return json.load(file)


def load_telecom_fixture() -> dict:
    with TELECOM_FIXTURE_PATH.open("r", encoding="utf-8") as file:
        return json.load(file)


def load_gas_natural_fixture() -> dict:
    with GAS_NATURAL_FIXTURE_PATH.open("r", encoding="utf-8") as file:
        return json.load(file)


def build_valid_context() -> dict:
    context = build_memorial_eletrico_v1_context(load_fixture())
    item_keys = [
        "cpct",
        "cftv",
        "alarme_patrimonial",
        "sonorizacao",
        "alarme_incendio",
        "automacao",
    ]
    context["nao_inclusos"]["tem_itens"] = any(
        bool(context["nao_inclusos"].get(key, False)) for key in item_keys
    )
    validate_memorial_eletrico_v1_context(context)
    return context


def build_eletrico_context_without_generator() -> dict:
    context = load_fixture()
    context["gerador"] = {
        "tem_gerador": False,
        "qtd": 0,
        "potencia_kva": 0,
        "tipo_atendimento": "condominio",
        "circuitos_atendidos": None,
    }
    built = build_memorial_eletrico_v1_context(context)
    validate_memorial_eletrico_v1_context(built)
    return built


def build_valid_telecom_context() -> dict:
    context = load_telecom_fixture()
    validate_memorial_telecom_v1_context(context)
    return context


def build_valid_gas_natural_context() -> dict:
    context = load_gas_natural_fixture()
    validate_memorial_gas_natural_v1_context(context)
    return context


def build_valid_glp_v2_context() -> dict:
    context = {
        "documento": {"data_atual": "13/05/2026"},
        "obra": {
            "numero_cadastro": "12345/2026",
            "construtora": "MGA CONSTRUCOES",
            "nome": "MAKAI",
            "localizacao": "Rua Exemplo, 100",
            "tipo_edificacao": "Residencial multifamiliar",
            "tipologia": "Subsolo, terreo, pavimentos tipo e cobertura",
            "qtd_apartamentos": {
                "valor": 29,
                "confianca": "high",
                "fonte_evidencia": [
                    {
                        "texto": "29 apartamentos",
                        "regra": "fixture_glp_v2_renderer",
                        "confianca": "high",
                    }
                ],
            },
            "qtd_lojas": 0,
            "qtd_restaurantes": 0,
        },
        "tanques": {
            "quantidade": 1,
            "tipo": "P-190",
            "capacidade_kg": 190,
            "qtd_abrigos": 1,
            "qtd_recipientes": 2,
            "fonte_evidencia": [
                {
                    "texto": "Abrigo de gas P-190",
                    "regra": "fixture_glp_v2_renderer",
                    "confianca": "high",
                }
            ],
            "conflitos": [],
        },
        "abastecimento": {"pavimento": "terreo"},
        "dimensionamento": {
            "qtd_fogao": 35,
            "qtd_aquecedor": 0,
            "qtd_churrasqueira": 35,
            "qtd_outros": 0,
        },
        "pontos_utilizacao": {
            "fogao": 35,
            "churrasqueira": 35,
            "aquecedor": 0,
            "outros": 0,
            "total_extraido": 70,
            "total_calculado": 70,
            "fontes_evidencia": [
                {
                    "texto": "35 fogoes e 35 churrasqueiras",
                    "regra": "fixture_glp_v2_renderer",
                    "confianca": "high",
                }
            ],
            "conflitos": [],
        },
        "diametros": {
            "tubulacao_principal": {
                "valor": 1.25,
                "unidade": "in",
                "valor_formatado": '1 1/4"',
                "valor_original": '1 1/4"',
            },
            "valvula_esfera": {
                "valor": 1.25,
                "unidade": "in",
                "valor_formatado": '1 1/4"',
                "valor_original": '1 1/4"',
                "inferido": True,
            },
        },
        "ramal": {
            "primario_material": "aco carbono SCH 40",
            "primario_pavimento": "terreo",
        },
        "numero": {"prancha": "01/04"},
        "teto_ou_piso": "piso",
        "context_version": "glp_v2",
        "template_version": "glp_v2",
    }
    validate_memorial_glp_v2_context(context)
    return context


class MemorialRendererTests(unittest.TestCase):
    @unittest.skipUnless(
        has_docx_render_dependencies(),
        "python-docx e docxtpl nao estao instalados no ambiente",
    )
    def test_render_memorial_eletrico_v1_generates_docx_without_template_tokens(self) -> None:
        tmp_dir = ROOT / "tests" / "output"
        output_path = tmp_dir / "renderer_test_output.docx"

        render_memorial_eletrico_v1(build_valid_context(), output_path)

        self.assertTrue(output_path.exists())
        text = inspect_docx_text(output_path)
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)
        self.assertIn("GRUPO GERADOR", text)

    @unittest.skipUnless(
        has_docx_render_dependencies(),
        "python-docx e docxtpl nao estao instalados no ambiente",
    )
    def test_render_memorial_eletrico_v1_omits_generator_section_when_absent(self) -> None:
        tmp_dir = ROOT / "tests" / "output"
        output_path = tmp_dir / "renderer_test_output_eletrico_sem_gerador.docx"

        render_memorial_eletrico_v1(build_eletrico_context_without_generator(), output_path)

        text = inspect_docx_text(output_path)
        self.assertNotIn("GRUPO GERADOR", text)
        self.assertNotIn("grupo geradores a diesel", text)

    def test_eletrico_context_keeps_mt_tension_empty_when_missing(self) -> None:
        context = build_memorial_eletrico_v1_context({"energia": {"tem_subestacao": False}})

        self.assertEqual(context["mt"], {})

    @unittest.skipUnless(
        has_docx_render_dependencies(),
        "python-docx e docxtpl nao estao instalados no ambiente",
    )
    def test_render_memorial_eletrico_v1_leaves_mt_section_blank_when_missing(self) -> None:
        tmp_dir = ROOT / "tests" / "output"
        output_path = tmp_dir / "renderer_test_output_eletrico_mt_blank.docx"
        context = build_valid_context()
        context["mt"] = {}

        render_memorial_eletrico_v1(context, output_path)

        text = inspect_docx_text(output_path)
        self.assertIn("circuito de  kV", text)
        self.assertNotIn("13.8", text)
        self.assertNotIn("13,8", text)
        self.assertNotIn("circuito de 13.8 kV", text)
        self.assertNotIn("circuito de 13,8 kV", text)

    @unittest.skipUnless(
        has_docx_render_dependencies(),
        "python-docx e docxtpl nao estao instalados no ambiente",
    )
    def test_render_memorial_telecom_v1_generates_docx_without_template_tokens(self) -> None:
        tmp_dir = ROOT / "tests" / "output"
        output_path = tmp_dir / "renderer_test_output_telecom.docx"

        render_memorial_telecom_v1(build_valid_telecom_context(), output_path)

        self.assertTrue(output_path.exists())
        text = inspect_docx_text(output_path)
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    @unittest.skipUnless(
        has_docx_render_dependencies(),
        "python-docx e docxtpl nao estao instalados no ambiente",
    )
    def test_render_memorial_gas_natural_v1_generates_docx_without_template_tokens(self) -> None:
        tmp_dir = ROOT / "tests" / "output"
        output_path = tmp_dir / "renderer_test_output_gas_natural.docx"

        render_memorial_gas_natural_v1(build_valid_gas_natural_context(), output_path)

        self.assertTrue(output_path.exists())
        text = inspect_docx_text(output_path)
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    @unittest.skipUnless(
        has_docx_render_dependencies(),
        "python-docx e docxtpl nao estao instalados no ambiente",
    )
    def test_render_memorial_gas_natural_v1_renders_valvula_esfera_diametro(self) -> None:
        tmp_dir = ROOT / "tests" / "output"
        output_path = tmp_dir / "renderer_test_output_gas_natural_valvula.docx"

        render_memorial_gas_natural_v1(build_valid_gas_natural_context(), output_path)

        text = inspect_docx_text(output_path)
        self.assertIn("32 mm", text)

    @unittest.skipUnless(
        has_docx_render_dependencies(),
        "python-docx e docxtpl nao estao instalados no ambiente",
    )
    def test_render_memorial_glp_v2_preserves_inch_diameters_and_schema_fields(self) -> None:
        tmp_dir = ROOT / "tests" / "output"
        output_path = tmp_dir / "renderer_test_output_glp_v2.docx"

        render_memorial_glp_v2(build_valid_glp_v2_context(), output_path)

        self.assertTrue(output_path.exists())
        text = inspect_docx_text(output_path)
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)
        self.assertNotIn("1.25mm", text)
        self.assertNotIn('1 1/4"mm', text)
        self.assertIn("MEMORIAL DESCRITIVO", text)
        self.assertIn("INSTALAÇÕES PREDIAIS DE GLP", text)
        self.assertIn('terá diâmetro de 1 1/4"', text)
        self.assertIn('válvula de esfera de 1 1/4"', text)
        self.assertIn("Fogão 04 bocas: 7.000 Kcal/h (35 unidades)", text)
        self.assertIn("Churrasqueira 04 bocas: 7.000 Kcal/h (35", text)
        self.assertIn("Totalizando 70 pontos de utilização", text)

    def test_glp_v2_template_uses_schema_v2_placeholders_not_legacy_v1_paths(self) -> None:
        from docx import Document

        template_text = "\n".join(
            paragraph.text for paragraph in Document(ROOT / "templates" / "glp" / "v2" / "template.docx").paragraphs
        )

        self.assertIn("{{ obra.qtd_apartamentos.valor }}", template_text)
        self.assertIn("{{ tanques.qtd_abrigos }}", template_text)
        self.assertNotIn("{{ tanques.quantidade }}", template_text)
        self.assertIn("{{ pontos_utilizacao.fogao }}", template_text)
        self.assertIn("{{ pontos_utilizacao.total_calculado }}", template_text)
        self.assertIn("{{ diametros.tubulacao_principal.valor_formatado }}", template_text)
        self.assertIn("{{ diametros.valvula_esfera.valor_formatado }}", template_text)
        self.assertNotIn("{{ obra.qtd_apartamentos }}", template_text)
        self.assertNotIn("{{ abastecimento.qtd_tanques }}", template_text)
        self.assertNotIn("{{ soma.qtd_pontos_de_utilizacao }}", template_text)
        self.assertNotIn("{{ ramal.primario_diametro }}", template_text)
        self.assertNotIn("xxxx", template_text)

    def test_assert_no_jinja_left_raises_structured_error(self) -> None:
        with self.assertRaises(MemorialRenderError) as error_info:
            assert_no_jinja_left("trecho com {{ placeholder }}")

        self.assertIn("Jinja", str(error_info.exception))

    def test_assert_no_internal_markers_left_raises_structured_error(self) -> None:
        with self.assertRaises(MemorialRenderError) as error_info:
            assert_no_internal_markers_left("Texto Fixo do template")

        self.assertIn("template", str(error_info.exception))


if __name__ == "__main__":
    unittest.main()
